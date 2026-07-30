"""High-risk branch coverage for non-standard quote mapping."""

from __future__ import annotations

import sys
from datetime import date
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook
from werkzeug.datastructures import FileStorage

import procurement_store
from services import (
    procurement_project_service,
    quote_mapping_service,
)


def _project_with_mapping_job(source):
    project_id = procurement_project_service.create_project(
        {
            'project_no': 'CG-MAPPING-TEST',
            'project_name': '报价映射分支测试',
        }
    )
    first_item = procurement_project_service.add_item(
        project_id,
        {
            'item_name': '结构件A',
            'spec_model': 'A-01',
            'quantity': '2',
            'unit': '件',
        },
    )
    second_item = procurement_project_service.add_item(
        project_id,
        {
            'item_name': '结构件B',
            'spec_model': 'B-02',
            'quantity': '3',
            'unit': '套',
        },
    )
    supplier_id = procurement_project_service.add_supplier(
        project_id,
        {'supplier_name': '映射供应商'},
    )
    job_id = procurement_store.create_mapping_job(
        {
            'project_id': project_id,
            'supplier_id': supplier_id,
            'quote_round': 1,
            'source_type': 'excel',
            'original_name': 'mapping.xlsx',
            'relative_path': 'mapping.xlsx',
            'file_sha256': 'mapping-hash',
            'source': source,
        }
    )
    return project_id, supplier_id, first_item, second_item, job_id


def test_mapping_extractors_bound_values_and_close_resources(
    tmp_path,
    monkeypatch,
):
    excel_path = tmp_path / 'mapping.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '报价'
    sheet.append(['日期', '内容'])
    sheet.append([date(2026, 7, 29), 'x' * 3000])
    workbook.save(excel_path)
    workbook.close()

    tables, diagnostics = quote_mapping_service._extract_excel(
        excel_path
    )
    assert diagnostics == []
    assert tables[0]['rows'][1][0] == '2026-07-29'
    assert len(tables[0]['rows'][1][1]) == 2000

    word_path = tmp_path / 'mapping.docx'
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = '物资'
    table.rows[0].cells[1].text = '单价'
    document.save(word_path)
    word_tables, _ = quote_mapping_service._extract_word(
        word_path
    )
    assert word_tables == [
        {'name': '表格1', 'rows': [['物资', '单价']]}
    ]

    class FakePage:
        def extract_tables(self):
            return [
                [
                    ['物资', '单价'],
                    None,
                    ['结构件A', None],
                ]
            ]

    class FakePdf:
        pages = [FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    monkeypatch.setitem(
        sys.modules,
        'pdfplumber',
        SimpleNamespace(open=lambda _path: FakePdf()),
    )
    pdf_tables, pdf_diagnostics = (
        quote_mapping_service._extract_pdf('mapping.pdf')
    )
    assert pdf_diagnostics == []
    assert pdf_tables[0]['rows'][1] == ['结构件A', '']

    captured = []
    queue = SimpleNamespace(put=captured.append)
    monkeypatch.setattr(
        quote_mapping_service,
        '_extract_excel',
        lambda _path: ([{'name': 'x', 'rows': []}], []),
    )
    quote_mapping_service._mapping_extract_worker(
        'excel',
        'ignored',
        queue,
    )
    assert captured[0][0] == 'ok'
    with pytest.raises(ValueError, match='不支持'):
        quote_mapping_service._mapping_extract_worker(
            'unknown',
            'ignored',
            queue,
        )


def test_create_mapping_job_rejects_invalid_scope_and_cleans_file(
    app,
    tmp_path,
    monkeypatch,
):
    project_id, supplier_id, *_ = _project_with_mapping_job(
        {'tables': [{'name': 'existing', 'rows': []}]}
    )
    upload = FileStorage(filename='quote.txt')
    with pytest.raises(ValueError, match='仅支持'):
        quote_mapping_service.create_mapping_job(
            project_id,
            supplier_id,
            1,
            upload,
        )
    with pytest.raises(ValueError, match='不存在'):
        quote_mapping_service.create_mapping_job(
            project_id + 999,
            supplier_id,
            1,
            FileStorage(filename='quote.xlsx'),
        )

    staged = tmp_path / 'quote.pdf'
    staged.write_bytes(b'%PDF-1.4')
    monkeypatch.setattr(
        quote_mapping_service.procurement_file_service,
        'save_upload',
        lambda *_args: {
            'absolute_path': str(staged),
            'relative_path': 'quote.pdf',
            'original_name': 'quote.pdf',
            'sha256': 'hash',
            'size_bytes': staged.stat().st_size,
        },
    )
    monkeypatch.setattr(
        quote_mapping_service,
        'run_isolated_worker',
        lambda *_args, **_kwargs: (
            [],
            ['没有可映射表格'],
        ),
    )
    with pytest.raises(ValueError, match='没有可映射表格'):
        quote_mapping_service.create_mapping_job(
            project_id,
            supplier_id,
            1,
            FileStorage(filename='quote.pdf'),
        )
    assert not staged.exists()


@pytest.mark.parametrize(
    ('form', 'message'),
    [
        ({'table_name': 'missing'}, '请选择有效'),
        (
            {'table_name': '报价', 'header_row': 'bad'},
            '表头行号无效',
        ),
        (
            {'table_name': '报价', 'header_row': '99'},
            '表头行号超出',
        ),
        (
            {'table_name': '报价', 'header_row': '1'},
            '至少映射',
        ),
        (
            {
                'table_name': '报价',
                'header_row': '1',
                'map_item_name': '0',
            },
            '必须映射单价',
        ),
        (
            {
                'table_name': '报价',
                'header_row': '1',
                'map_item_name': 'not-a-column',
                'map_unit_price': '1',
            },
            '物资名称列映射无效',
        ),
        (
            {
                'table_name': '报价',
                'header_row': '1',
                'map_item_name': '0',
                'map_unit_price': '99',
            },
            '单价列映射超出表头范围',
        ),
    ],
)
def test_map_to_import_job_rejects_invalid_mapping(
    app,
    form,
    message,
):
    _, _, _, _, job_id = _project_with_mapping_job(
        {
            'tables': [
                {
                    'name': '报价',
                    'rows': [
                        ['物资', '单价'],
                        ['结构件A', '10'],
                    ],
                }
            ]
        }
    )
    with pytest.raises(ValueError, match=message):
        quote_mapping_service.map_to_import_job(job_id, form)


def test_map_to_import_job_records_errors_warnings_and_redacts_csrf(
    app,
):
    _, _, _, _, job_id = _project_with_mapping_job(
        {
            'tables': [
                {
                    'name': '报价',
                    'rows': [
                        [
                            '明细ID',
                            '行号',
                            '名称',
                            '型号',
                            '数量',
                            '单价',
                            '金额',
                        ],
                        [
                            'bad',
                            'bad',
                            '结构件A',
                            'A-01',
                            '2',
                            '10',
                            '999',
                        ],
                        [
                            '',
                            '1',
                            '',
                            '',
                            '2',
                            '11',
                            '',
                        ],
                        [
                            '',
                            '',
                            '不存在的物资',
                            '',
                            '1',
                            '5',
                            '',
                        ],
                    ],
                }
            ],
            'diagnostics': ['源文件提示'],
            'size_bytes': 123,
        }
    )
    import_job_id = quote_mapping_service.map_to_import_job(
        job_id,
        {
            'table_name': '报价',
            'header_row': '1',
            'map_project_item_id': '0',
            'map_line_no': '1',
            'map_item_name': '2',
            'map_spec_model': '3',
            'map_quantity': '4',
            'map_unit_price': '5',
            'map_amount': '6',
            'tax_rate': 'Infinity',
            'csrf_token': 'must-not-persist',
        },
    )

    import_job = procurement_store.get_import_job(import_job_id)
    assert len(import_job['payload']['items']) == 1
    assert import_job['payload']['items'][0][
        'amount_minor'
    ] == 2000
    assert any(
        '金额与数量×单价不一致' in item
        for item in import_job['warnings']
    )
    assert any(
        '无法匹配项目明细' in item
        for item in import_job['warnings']
    )
    assert any(
        '重复匹配项目明细' in item
        for item in import_job['errors']
    )
    assert '税率格式无效' in import_job['errors']
    mapping_job = procurement_store.get_mapping_job(job_id)
    assert mapping_job['status'] == 'invalid'
    assert 'csrf_token' not in mapping_job['metadata']


def test_map_to_import_job_supports_line_matching_and_complete_payload(
    app,
):
    _, _, _, _, job_id = _project_with_mapping_job(
        {
            'tables': [
                {
                    'name': '报价',
                    'rows': [
                        ['行号', '数量', '单位', '单价'],
                        ['1', '2', '件', '10.50'],
                        ['2', '3', '套', '20'],
                    ],
                }
            ],
            'size_bytes': 50,
        }
    )
    import_job_id = quote_mapping_service.map_to_import_job(
        job_id,
        {
            'table_name': '报价',
            'header_row': '1',
            'map_line_no': '0',
            'map_quantity': '1',
            'map_unit': '2',
            'map_unit_price': '3',
            'tax_rate': '13',
            'price_basis': 'tax_inclusive',
        },
    )

    import_job = procurement_store.get_import_job(import_job_id)
    assert import_job['errors'] == []
    assert import_job['warnings'] == []
    assert len(import_job['payload']['items']) == 2
    assert import_job['payload']['header']['tax_rate_bps'] == 1300
    assert (
        import_job['payload']['header']['total_amount_minor']
        == 8100
    )
    assert (
        procurement_store.get_mapping_job(job_id)['status']
        == 'parsed'
    )


def test_map_to_import_job_rejects_tax_precision_beyond_basis_points(
    app,
):
    _, _, _, _, job_id = _project_with_mapping_job(
        {
            'tables': [
                {
                    'name': '报价',
                    'rows': [
                        ['行号', '单价'],
                        ['1', '10'],
                    ],
                }
            ]
        }
    )

    import_job_id = quote_mapping_service.map_to_import_job(
        job_id,
        {
            'table_name': '报价',
            'header_row': '1',
            'map_line_no': '0',
            'map_unit_price': '1',
            'tax_rate': '13.333',
        },
    )

    import_job = procurement_store.get_import_job(import_job_id)
    assert '税率最多保留两位小数' in import_job['errors']
    assert import_job['payload']['header']['tax_rate_bps'] is None
    assert procurement_store.get_mapping_job(job_id)['status'] == 'invalid'


def test_map_to_import_job_rejects_missing_job(app):
    with pytest.raises(ValueError, match='字段映射任务不存在'):
        quote_mapping_service.map_to_import_job(
            999999,
            {},
        )
